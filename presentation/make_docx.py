#!/usr/bin/env python3
"""Build the report as .docx from the markdown, embedding charts at section anchors."""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE=os.path.dirname(__file__)
ROOT=os.path.dirname(HERE)
MD=os.path.join(ROOT,"report","China_AI_Stack_Report.md")
GLOSS=os.path.join(ROOT,"report","GLOSSARY.md")
CH=os.path.join(HERE,"charts")
OUT=os.path.join(ROOT,"report","China_AI_Stack_Report.docx")

INK=RGBColor(0x18,0x1A,0x1F); ACCENT=RGBColor(0xB0,0x2A,0x1B); MUTED=RGBColor(0x56,0x5C,0x67)
STEEL=RGBColor(0x2C,0x6E,0x7A); FAINT=RGBColor(0x7A,0x82,0x8E)
BODY="Georgia"; SANS="Calibri"; MONO="Consolas"

# chart insert: after a heading whose text startswith key -> (image, width_in, caption)
CHARTS={
 "2. The macro frame":("02_supply_demand.png",5.7,"Figure 1 — The supply gap that built the market. Sources: REF-SANDS-01 (C), REF-BW-01 (D)."),
 "3. The stack, layer by layer":("01_value_chain_map.png",6.4,"Figure 2 — Value-chain map: China's AI stack, bottom-up, with capability vs. the global frontier."),
 "3.0 Energy & Power":("04_power.png",5.7,"Figure 3 — Electricity generation, China vs. US, 2024. Sources: WEB-001 (C)."),
 "3.1 Semiconductor equipment":("05_wfe_selfsuff.png",5.9,"Figure 4 — L1 self-sufficiency by segment. Sources: WEB-009/011/019 (C→D)."),
 "3.2 Fabrication":("03_yield_duv.png",6.2,"Figure 5 — 7nm-class yield and the DUV multipatterning penalty. Sources: WEB-014 (C)."),
 "3.7 Foundation models":("06_token_share.png",5.9,"Figure 6 — Chinese-model share of OpenRouter tokens, 2026 (C)."),
 "4.4 Capital markets":("07_ipo_pops.png",5.9,"Figure 7 — 2025–26 China AI-stack IPOs, debut price moves. Sources: WEB (B/C)."),
 "5.2 China-vs-US scorecard":("08_scorecard.png",5.2,"Figure 8 — Scorecard: capability gap × foreign investability, by layer."),
 "II.1 The demand split":("13_demand_spectrum.png",6.2,"Figure II-1 — Global-demand-led vs domestic-substitution: where each name's demand comes from (directional, C/D)."),
 "II.3 Model & app economics":("10_model_econ.png",5.9,"Figure II-2 — Model-lab economics: market share ≠ revenue (Zhipu & MiniMax FY2025 filed, B)."),
 "II.4 Talent & R&D":("12_talent.png",5.9,"Figure II-3 — China's AI research output vs its historical talent export (C)."),
 "II.6 Index exposure":("11_index_weight.png",5.9,"Figure II-4 — AI-stack weight inside major China indices (named leaders = a floor; C)."),
}

def set_font(run,name=BODY,size=10.5,color=INK,bold=False,italic=False):
    run.font.name=name; run.font.size=Pt(size); run.font.color.rgb=color
    run.bold=bold; run.italic=italic
    r=run._element.rPr.rFonts; r.set(qn('w:eastAsia'),name)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)

def add_inline(p,text,base=BODY,size=10.5,color=INK):
    # split on **bold**, *italic* and `code`
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)',text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**"):
            set_font(p.add_run(tok[2:-2]),base,size,color,bold=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok)>2:
            set_font(p.add_run(tok[1:-1]),base,size,color,italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            set_font(p.add_run(tok[1:-1]),MONO,size-0.5,STEEL)
        else:
            set_font(p.add_run(tok),base,size,color)

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
for m in ("top_margin","bottom_margin"): setattr(sec,m,Inches(0.9))
for m in ("left_margin","right_margin"): setattr(sec,m,Inches(1.0))

def heading(text,level):
    p=doc.add_paragraph(); p.space_before=Pt(14 if level<=1 else 10); p.paragraph_format.space_after=Pt(4)
    sizes={0:22,1:16,2:12.5,3:11}; cols={0:ACCENT,1:ACCENT,2:INK,3:STEEL}
    add=p.add_run(text); set_font(add,SANS,sizes[level],cols[level],bold=True)
    return p

def caption(text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(12)
    set_font(p.add_run(text),SANS,8.5,FAINT,italic=True)

def insert_chart_for(htext):
    for key,(img,w,cap) in CHARTS.items():
        if htext.startswith(key):
            path=os.path.join(CH,img)
            if os.path.exists(path):
                pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                pp.paragraph_format.space_before=Pt(8); pp.paragraph_format.space_after=Pt(2)
                pp.add_run().add_picture(path,width=Inches(w))
                cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                set_font(cp.add_run(cap),SANS,8.5,FAINT,italic=True)
            break

# ---- title block ----
tp=doc.add_paragraph(); set_font(tp.add_run("The Stack China Is Building"),SANS,30,ACCENT,bold=True)
sp=doc.add_paragraph(); set_font(sp.add_run("A Value-Chain Map of China's Indigenous AI Ecosystem"),SANS,15,INK,bold=True)
mp=doc.add_paragraph(); set_font(mp.add_run("Flagship research report  ·  Data as-of mid-2026  ·  For financial-market experts, economists & industry analysts"),SANS,10,MUTED)
# hero map on cover
hp=doc.add_paragraph(); hp.alignment=WD_ALIGN_PARAGRAPH.CENTER; hp.paragraph_format.space_before=Pt(10)
hp.add_run().add_picture(os.path.join(CH,"01_value_chain_map.png"),width=Inches(6.4))
doc.add_page_break()

# ---- parse markdown ----
lines=open(MD).read().splitlines()
if os.path.exists(GLOSS):
    lines += ["","<<<PAGEBREAK>>>",""] + open(GLOSS).read().splitlines()
i=0; n=len(lines); started=False   # skip md's own title block (duplicated on cover); keep callout
while i<n:
    ln=lines[i]
    st=ln.strip()
    if not st:
        i+=1; continue
    if st=="<<<PAGEBREAK>>>":
        doc.add_page_break(); started=True; i+=1; continue
    if not started:
        if st.startswith("## "):
            started=True            # real content begins
        elif st.startswith("> "):
            pass                    # allow the confidence callout through
        else:
            i+=1; continue          # skip duplicate title / subtitle / italic meta
    if st.startswith("> "):   # callout
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.18)
        p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6)
        pPr=p._p.get_or_add_pPr(); bdr=OxmlElement('w:pBdr'); left=OxmlElement('w:left')
        left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'8'); left.set(qn('w:color'),'B02A1B')
        bdr.append(left); pPr.append(bdr)
        add_inline(p,st[2:],base=BODY,size=10,color=MUTED)
        i+=1; continue
    if st.startswith("#"):
        m=re.match(r'(#+)\s+(.*)',st); lvl=len(m.group(1))-1; txt=m.group(2)
        txt=re.sub(r'\*','',txt)
        txt=re.sub(r'\s*‹[^›]*›','',txt).strip()
        heading(txt,min(lvl,3))
        insert_chart_for(txt)
        i+=1; continue
    if st.startswith("---"):
        i+=1; continue
    if st.startswith("|") and i+1<n and re.match(r'^\s*\|[\s:|-]+\|\s*$',lines[i+1]):
        # table
        rows=[]
        while i<n and lines[i].strip().startswith("|"):
            if not re.match(r'^\s*\|[\s:|-]+\|\s*$',lines[i]):
                cells=[c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
            i+=1
        if rows:
            ncol=len(rows[0]); t=doc.add_table(rows=0,cols=ncol); t.alignment=WD_TABLE_ALIGNMENT.CENTER
            t.style="Table Grid"
            for ri,r in enumerate(rows):
                cells=t.add_row().cells
                for ci in range(ncol):
                    val=r[ci] if ci<len(r) else ""
                    cell=cells[ci]; cell.paragraphs[0].text=""
                    add_inline(cell.paragraphs[0],val,base=SANS,size=8.8,color=INK if ri else RGBColor(0xFF,0xFF,0xFF))
                    if ri==0: shade(cell,"B02A1B")
        continue
    if st.startswith("- ") or st.startswith("* "):
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2)
        add_inline(p,st[2:],base=BODY,size=10,color=INK)
        i+=1; continue
    # normal paragraph
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.25
    add_inline(p,st,base=BODY,size=10.5,color=INK)
    i+=1

doc.save(OUT)
print("saved",OUT, os.path.getsize(OUT),"bytes")
