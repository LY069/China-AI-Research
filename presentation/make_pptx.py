#!/usr/bin/env python3
"""Build a native, editable PowerPoint deck (16:9) mirroring the HTML deck,
reusing the chart PNGs in presentation/charts/."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE=os.path.dirname(__file__); CH=os.path.join(HERE,"charts")
OUT=os.path.join(os.path.dirname(HERE),"presentation","China_AI_Stack_Deck.pptx")

INK=RGBColor(0x18,0x1A,0x1F); ACCENT=RGBColor(0xC0,0x2A,0x1B); MUTED=RGBColor(0x56,0x5C,0x67)
STEEL=RGBColor(0x2C,0x6E,0x7A); COPPER=RGBColor(0xA9,0x6E,0x28); FAINT=RGBColor(0x8A,0x91,0x9C)
PAPER=RGBColor(0xEC,0xED,0xEF); WHITE=RGBColor(0xFF,0xFF,0xFF); GOOD=RGBColor(0x2C,0x7A,0x4B)
SANS="Calibri"; MONO="Consolas"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=prs.slide_width,prs.slide_height

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s,color=WHITE):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=color

def box(s,l,t,w,h):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; return tb,tf

def setrun(r,text,size,color=INK,bold=False,font=SANS,italic=False):
    r.text=text; f=r.font; f.size=Pt(size); f.color.rgb=color; f.bold=bold; f.name=font; f.italic=italic

def para(tf,text,size,color=INK,bold=False,font=SANS,space_after=6,bullet=False,align=PP_ALIGN.LEFT,first=False):
    p=tf.paragraphs[0] if first and tf.paragraphs[0].text=="" else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space_after)
    setrun(p.add_run(),text,size,color,bold,font); return p

def eyebrow(s,text,num):
    _,tf=box(s,0.6,0.42,11.5,0.4)
    p=tf.paragraphs[0]
    setrun(p.add_run(),text.upper()+"   ",11,ACCENT,True,MONO)
    setrun(p.add_run(),"· "+num,11,FAINT,False,MONO)

def title(s,text,y=0.95,size=30,color=INK):
    _,tf=box(s,0.6,y,12.1,1.2); para(tf,text,size,color,True,SANS,first=True)

def accentbar(s):
    from pptx.enum.shapes import MSO_SHAPE
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.6),Inches(0.34),Inches(0.55),Inches(0.06))
    sh.fill.solid(); sh.fill.fore_color.rgb=ACCENT; sh.line.fill.background()

def img(s,path,l,t,w):
    if os.path.exists(path):
        s.shapes.add_picture(path,Inches(l),Inches(t),width=Inches(w))

def ch(name): return os.path.join(CH,name)

# ---------- 1 TITLE ----------
s=slide(); bg(s,PAPER); accentbar(s)
_,tf=box(s,0.6,0.34,11,0.4); para(tf,"INDIGENOUS AI ECOSYSTEM   ·   DATA AS-OF MID-2026",11,MUTED,True,MONO,first=True)
_,tf=box(s,0.6,2.1,12.1,2.2); para(tf,"The Stack China Is Building",46,ACCENT,True,SANS,first=True)
_,tf=box(s,0.6,3.9,11.8,1.6)
para(tf,"A value-chain map of China's indigenous AI ecosystem — where capability, capital and captive demand meet, and where a professional investor can (and cannot) get exposure.",18,INK,False,SANS,first=True)
_,tf=box(s,0.6,6.4,12,0.6); para(tf,"Prepared for financial-market experts, economists & industry analysts · Every figure carries a source ID + A–D confidence grade",11,FAINT,False,MONO,first=True)

# helper for a standard content slide with eyebrow+title+optional image+bullets
def content(eb,num,ttl,bullets=None,image=None,img_w=8.6,img_l=2.35,img_t=2.15,note=None,title_size=28,two_col_bullets=None):
    s=slide(); bg(s,WHITE); accentbar(s); eyebrow(s,eb,num); title(s,ttl,size=title_size)
    y=2.05
    if image:
        img(s,ch(image),img_l,img_t,img_w); y=img_t+5.0
    if bullets:
        _,tf=box(s,0.6,2.0 if not image else 5.9,12.1,4.6 if not image else 1.4)
        for i,(txt,c,bold) in enumerate(bullets):
            para(tf,txt,15 if not image else 13,c,bold,SANS,space_after=7,first=(i==0))
    if note:
        _,tf=box(s,0.6,6.9,12.1,0.4); para(tf,note,9.5,FAINT,False,MONO,first=True)
    return s

B=lambda t,c=INK,bold=False:(t,c,bold)

# ---------- 2 THESIS ----------
content("The Thesis","02","One line, then the map",
    bullets=[
      B("China is building a PARALLEL, vertically-indigenized AI stack under export-control pressure — one to several years behind the US frontier at almost every layer, but with the direction, coordination and capital to create a distinct investable ecosystem whose value is migrating to hardware, power and embodied AI.",INK,True),
      B("Constrained:  lithography is the binding chokepoint · software (CUDA) the deepest gap (3–5+ yr) · compute rationed (~2.5 GW inference).",MUTED,False),
      B("Working:  a system-level bet (lose at the chip, win at the rack) · power, the one clear edge · a 'slow-bull' IPO wave as the 2026 catalyst.",MUTED,False),
    ])

# ---------- 3 DOCTRINES ----------
content("Why It's Its Own Universe","03","Two doctrines of AI",
    bullets=[
      B("United States — a race to superintelligence:  frontier labs treat a few months' lead as potentially permanent; effort concentrates on the frontier model.",INK,True),
      B("China — AI as infrastructure:  like electricity or water, to be diffused across the whole economy, with a pronounced tilt to the PHYSICAL world (manufacturing, robotics).",ACCENT,True),
      B("This sets where capital flows (systems & diffusion, not only frontier training), and why China tolerates a chip deficit while racing on rack-scale systems, power and embodied AI. Export controls reframed every dependency as a weaponizable vulnerability — so China indigenizes every choke point.",MUTED,False),
    ])

# ---------- 4 VALUE-CHAIN MAP (hero) ----------
content("The Map","04 · Hero","The stack, bottom-up — and China's position",
    image="01_value_chain_map.png",img_w=9.4,img_l=1.95,img_t=1.95,
    note="Highlighted layers (L1 equipment · L4 chips · L9 physical AI) are where the report concentrates.")

# ---------- 5 COMPUTE S/D ----------
content("Macro Frame","05","A low compute base — and prices are turning",
    image="02_supply_demand.png",img_w=8.2,img_l=2.55,img_t=2.2,
    note="After years of discounting, Alibaba/Baidu/Tencent RAISED AI-compute list prices in 2026 — compute, not demand, is the binding constraint. REF-BW-01/SANDS-01.")

# ---------- 6 LITHOGRAPHY ----------
content("L1 · The Binding Constraint","06","Without EUV, everything upstream is rationed",
    image="03_yield_duv.png",img_w=9.2,img_l=2.05,img_t=2.2,
    note="Indigenous EUV light source ~100–150W mid-2025 vs >250W needed for HVM — the '2028 goal' is a lab prototype, not parity. WEB-014/007.")

# ---------- 7 CHIP VS RACK ----------
content("L4 · L5 · The Central Battleground","07","Lose at the chip, win at the rack",
    bullets=[
      B("Device level:  Ascend 910C ≈ 60% of an H100 on inference (DeepSeek's own benchmark) — ~1–2 generations behind.",INK,True),
      B("System level:  CloudMatrix / Atlas 950 lash up to 8,192 NPUs into one logical machine (16 PB/s interconnect), competing with Nvidia's GB200 NVL72 at the rack — at ~4× the power.",INK,True),
      B("The deep moat is SOFTWARE:  CANN/MindSpore vs CUDA is a 3–5+ year gap, even as Huawei open-sources CANN and DeepSeek/Zhipu train without Nvidia.",ACCENT,True),
      B("Demand engine:  state-DC domestic-chip rules escalated 50% → full foreign-accelerator ban → ≥80%-domestic national grid. (Nvidia's China share reporting is unreconciled.)",MUTED,False),
    ])

# ---------- 8 POWER ----------
content("L0 · The Structural Edge","08","Power is the one unambiguous advantage",
    image="04_power.png",img_w=8.2,img_l=2.55,img_t=2.25,
    note="But it is NECESSARY, NOT SUFFICIENT — the payoff is gated on domestic silicon closing the compute gap (a D-grade inference). Cleanest proxy: CATL.")

# ---------- 9 MODELS ----------
content("L7 · L8 · Models & Diffusion","09","Commoditize intelligence — on purpose",
    image="06_token_share.png",img_w=8.4,img_l=2.45,img_t=2.25,
    note="Chinese open-weight models take up to ~55% of OpenRouter tokens. Strategy: make good-enough intelligence cheap infrastructure — pressuring Western labs' pricing.")

# ---------- 10 SCORECARD ----------
content("Synthesis","10","China-vs-US scorecard: capability × access",
    image="08_scorecard.png",img_w=7.4,img_l=2.95,img_t=1.95,
    note="Green = ahead / open · red = behind / closed. The purest exposure (L4 chips) is the least accessible to US capital.")

# ---------- 11 CAPITAL MARKETS ----------
content("Capital Markets","11","The 'slow-bull' IPO wave is the 2026 catalyst",
    image="07_ipo_pops.png",img_w=8.6,img_l=2.35,img_t=2.2,
    note="Montage's A/H premium inverted (H from a 44% discount to a premium). CXMT (~$4.3bn) & YMTC are the next, much larger tests.")

# ---------- 12 INVESTABILITY OVERLAY ----------
content("Decisive For Allocators","12","Investability ≠ listing",
    bullets=[
      B("Treasury's OISP now overlaps Entity-List flags: the highest-momentum GPU names trade freely in Hong Kong yet are effectively CLOSED to US capital.",ACCENT,True),
      B("Open · high-conviction:  Alibaba, Tencent, Montage, CATL, Baidu.",GOOD,True),
      B("Access-constrained:  NAURA, AMEC, ACM, SMIC, Hua Hong, Cambricon (HK lines / Connect where eligible).",COPPER,True),
      B("Closed / pipeline:  Moore Threads, MetaX, Biren, Kunlunxin, Huawei, ByteDance; CXMT/YMTC (IPO).",INK,False),
      B("Cleaner embodied-AI access:  WeRide, Pony.ai, Momenta, UBTech, Horizon Robotics.",MUTED,False),
    ])

# ---------- 13 SCENARIOS ----------
content("Scenarios & Risk","13","2027 / 2030 paths",
    bullets=[
      B("Base:  uneven indigenization — foundry/memory/interconnect/power progress, EUV lags; competitive rack-scale systems & open models; equities compound off a low base, high dispersion.",INK,True),
      B("Bull:  an EUV or HBM breakthrough closes the compute gap; the power edge converts; champions re-rate toward global peers.",GOOD,True),
      B("Bear:  EUV stalls, yields stay low, mandate outruns capability; a Taiwan / rare-earth flashpoint or OISP tightening strands foreign capital; policy-inflated multiples de-rate.",ACCENT,True),
      B("Key gating risk:  technology failure at EUV / HBM.  Most underappreciated practical risk:  the Entity-List + OISP overlay stranding US capital.",MUTED,True),
    ])

# ---------- 14 PART II — DEMAND SPLIT ----------
content("Part II · Demand","14","Not all 'China-AI' is a bet on China",
    image="13_demand_spectrum.png",img_w=9.6,img_l=1.85,img_t=2.05,
    note="The diversifying, low-correlation thesis applies ONLY to the domestic-substitution bucket. Global-demand names price off the world AI cycle.")

# ---------- 15 PART II — MODEL ECONOMICS ----------
content("Part II · Economics","15","Market share ≠ revenue",
    image="10_model_econ.png",img_w=8.6,img_l=2.35,img_t=2.2,
    note="Open models win token share but earn little directly. DeepSeek/Moonshot/ByteDance revenue not reliably disclosed — declined to state. Baidu 6-K: AI-Apps revenue ~flat YoY.")

# ---------- 16 PART II — GLOBAL ACCESS ----------
content("Part II · Access","16","The US is the outlier, not the norm",
    bullets=[
      B("EU · UK · Japan · Korea · Singapore · Gulf:  NO domicile-specific legal bar on most listed China-AI names (incl. Entity-Listed GPU IPOs) — access gated by market mechanics (Stock Connect, QFII, STAR pro-only), not law.",GOOD,True),
      B("United States:  the exception (Entity List + OISP; small passive-index holdings exempt).",ACCENT,True),
      B("Taiwan:  the one symmetric legal barrier (restricts its own capital into mainland tech). Japan's GPIF voluntarily excludes China A (2025–30).",INK,False),
      B("Swing factor:  a single-sourced report that Beijing may curb foreign access to its AI firms — the first China-side capital wall if enacted (monitor).",MUTED,True),
    ])

# ---------- 17 PART II — INDEX EXPOSURE ----------
content("Part II · Index","17","You may already own the stack",
    image="11_index_weight.png",img_w=8.6,img_l=2.35,img_t=2.2,
    note="Tencent+Alibaba ≈ 23.5% of MSCI China; SMIC+Tencent ≈ 17.9% of HSTECH; Innolight the top CSI-300 weight. Passive China exposure already carries heavy AI-stack beta.")

# ---------- 18 PART II — TALENT ----------
content("Part II · Talent","18","Leads on research output; has exported its talent",
    image="12_talent.png",img_w=8.6,img_l=2.35,img_t=2.2,
    note="China: 69.7% of AI patent filings, 38% of elite researchers by origin — but historically 72% went to US labs. Migration into the US fell ~89% since 2017 (reflow signal).")

# ---------- 19 CLOSE ----------
s=slide(); bg(s,PAPER); accentbar(s); eyebrow(s,"Deliverables & Method","19"); title(s,"What backs this deck")
_,tf=box(s,0.6,2.1,12.1,3.6)
for i,(t,c,b) in enumerate([
    ("Report — long-form value-chain map + Part II (demand, economics, access) + glossary; 12 charts.",INK,True),
    ("Database — 43 companies, metrics, policy/funding, model & chip catalogs, company_reference (ISIN/GICS/index), 100+ sources.",INK,True),
    ("Evidence — 12 analyst pod memos across two rounds, plus an adversarial red-team review.",INK,True),
    ("Every figure is source-tagged and A–D confidence-graded. Key open questions are documented, not smoothed: Nvidia's true China share, CXMT/YMTC valuations, indigenous-EUV timeline, model-lab revenue.",MUTED,False),
    ("Reference scaffolding: Bridgewater ×3, Sands Capital, Exponential View. Analysis, not investment advice.",FAINT,False),
]):
    para(tf,t,15,c,b,SANS,space_after=10,first=(i==0))

prs.save(OUT)
print("saved",OUT,os.path.getsize(OUT),"bytes,",len(prs.slides.__iter__.__self__._sldIdLst),"slides")
